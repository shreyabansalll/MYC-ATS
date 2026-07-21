# pipeline/generator.py
# STAGE 7 — ATS-Compliant DOCX Generator
# Takes rewritten JSON from services/ai_client.py
# Builds a clean single-column ATS-friendly DOCX
# Library: python-docx
# Rules applied: R1-R15 (format), R16-R25 (sections) from ATS PDF

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
from config import OUTPUT_DIR


# ── Rank groups for page length control ──────────────────────────────
SINGLE_PAGE_RANKS = {
    'deck cadet', 'engine cadet', 'oiler', 'motorman',
    'fitter', 'ab seaman', 'able seaman', 'bosun', 'jwko'
}
TWO_PAGE_RANKS = {
    'third officer', 'second officer', 'third engineer',
    'fourth engineer', 'eto', 'second engineer'
}
# master, chief officer, chief engineer → 2 pages, more vessel entries allowed


def _as_list(value) -> list:
    """
    Defensive coercion for LLM output that doesn't follow the expected
    shape (e.g. a string instead of a list for skills/certifications/
    education/languages) — applied consistently across all list-shaped
    content fields rather than ad hoc per field.
    """
    if isinstance(value, list):
        return value
    return [value] if value else []


def split_into_sentences(text: str) -> list:
    """
    Splits text into sentences on '.', '!', '?' — but treats a period
    immediately following a single capital letter (e.g. 'M.T', 'M.V' —
    standard maritime vessel-prefix abbreviations for Motor Tanker/Motor
    Vessel) as not a sentence boundary. A naive split('.') here truncates
    summaries mid-abbreviation: confirmed on real generated output where
    'Sailed on M.T Example GRT 12345...' got cut to '...Sailed on M.T'
    once the 3-sentence cap kicked in after only two real sentences.
    """
    if not text:
        return []
    text = text.replace('!', '.').replace('?', '.')
    tokens = text.split('.')
    sentences = []
    current = ''
    for i, tok in enumerate(tokens):
        current += tok
        if i == len(tokens) - 1:
            break
        last_word = current.split()[-1] if current.split() else ''
        if len(last_word) == 1 and last_word.isupper():
            current += '.'   # part of an abbreviation like "M." — don't split
            continue
        sentences.append(current.strip())
        current = ''
    if current.strip():
        sentences.append(current.strip())
    return [s for s in sentences if s]


def enforce_length_limits(content: dict, rank: str) -> dict:
    rank_lower = rank.lower() if rank else ''
    is_junior = rank_lower in SINGLE_PAGE_RANKS

    # Skills: 15 for junior, 18 for senior
    max_skills = 15 if is_junior else 18
    content['skills'] = _as_list(content.get('skills', []))[:max_skills]

    # Summary: 3 sentences max always
    summary = content.get('summary', '')
    sentences = split_into_sentences(summary)
    content['summary'] = '. '.join(sentences[:3]) + ('.' if sentences else '')

    # Experience: limit entries and bullets per entry
    experience = _as_list(content.get('experience_bullets', []))
    max_entries = 3 if is_junior else 6
    experience = experience[:max_entries]
    for job in experience:
        if isinstance(job, dict):
            job['bullets'] = job.get('bullets', [])[:3]
    content['experience_bullets'] = experience

    # Certifications: cap at 8 lines
    content['certifications'] = _as_list(content.get('certifications', []))[:8]

    return content


def add_horizontal_rule(paragraph):
    """
    Adds a thin bottom border under a paragraph.
    Used under section headings — replaces table-based dividers (ATS violation).
    """
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '2E75B6')
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_section_heading(doc, text):
    """
    Adds a section heading: Arial 12pt bold ALL CAPS + blue underline rule.
    No tables, no text boxes — pure paragraph with border.
    """
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = 'Arial'
    run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    add_horizontal_rule(p)
    return p


def add_body_text(doc, text, bold=False, size=11):
    """
    Adds a plain body paragraph — Arial, 11pt, no special formatting.
    """
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Arial'
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_bullet(doc, text):
    """
    Adds a bullet point using Word's built-in List Bullet style.
    Never uses unicode bullets (ATS violation R11).
    """
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = 'Arial'
    p.paragraph_format.space_after = Pt(1)
    return p


def generate_docx(content: dict, job_id: str = 'output') -> str:
    """
    Main entry point.
    Takes rewritten content JSON from services/ai_client.py.
    Returns path to the generated DOCX file.

    Expected content keys:
    - name, contact (dict), summary, experience_bullets (list),
      skills (list), certifications (list), rank (str)
    """
    # Enforce length limits before building
    rank = content.get('rank', '')
    content = enforce_length_limits(content, rank)

    doc = Document()

    # ── Page Setup (R13, R14) ─────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1)
        section.right_margin  = Inches(1)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

    # Set default font
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)

    # ── CANDIDATE NAME (R16) ──────────────────────────────────────────
    name_p = doc.add_paragraph()
    name_run = name_p.add_run(content.get('name', 'Candidate Name'))
    name_run.bold = True
    name_run.font.size = Pt(18)
    name_run.font.name = 'Arial'
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_after = Pt(2)

    # ── CONTACT INFO — in body, never in header (R6, R16) ────────────
    contact = content.get('contact', {})
    contact_parts = []
    if contact.get('email'):    contact_parts.append(contact['email'])
    if contact.get('phone'):    contact_parts.append(contact['phone'])
    if contact.get('linkedin'): contact_parts.append(contact['linkedin'])
    if contact.get('github'):   contact_parts.append(contact['github'])

    contact_p = doc.add_paragraph()
    contact_run = contact_p.add_run(' | '.join(contact_parts))
    contact_run.font.size = Pt(10)
    contact_run.font.name = 'Arial'
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_after = Pt(8)

    # ── PROFESSIONAL SUMMARY (R19) ────────────────────────────────────
    if content.get('summary'):
        add_section_heading(doc, 'Professional Summary')
        add_body_text(doc, content['summary'])

    # ── CERTIFICATIONS (R23) ─────────────────────────────────────────
    # Placed immediately after the summary, ahead of Work Experience.
    # Maritime recruiters scan resumes in a strict priority order —
    # certificates first (the eligibility gate), then sea service, then
    # vessel type experience (see CLAUDE.md's Project Overview) — so the
    # #1 thing they check for shouldn't be buried mid-document. Section
    # detection in pipeline/scorer.py is presence-based, not
    # position-based, so this has no effect on the ATS score.
    certifications = _as_list(content.get('certifications', []))
    if certifications:
        add_section_heading(doc, 'Certifications')
        for cert in certifications:
            add_body_text(doc, str(cert))

    # ── WORK EXPERIENCE (R21, R22) ────────────────────────────────────
    experience = _as_list(content.get('experience_bullets', []))
    if experience:
        add_section_heading(doc, 'Work Experience')
        for job in experience:
            if not isinstance(job, dict):
                continue
            role_line = f"{job.get('role', '')}  —  {job.get('company', '')}  |  {job.get('dates', '')}"
            add_body_text(doc, role_line, bold=True)
            for bullet in job.get('bullets', []):
                add_bullet(doc, bullet)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ── SKILLS (R20) ──────────────────────────────────────────────────
    skills = _as_list(content.get('skills', []))
    if skills:
        add_section_heading(doc, 'Skills')
        skills_text = '  •  '.join(str(s) for s in skills)
        add_body_text(doc, skills_text)

    # ── EDUCATION ─────────────────────────────────────────────────────
    education = _as_list(content.get('education', []))
    if education:
        add_section_heading(doc, 'Education')
        for edu in education:
            add_body_text(doc, str(edu))

    # ── DOCUMENTS & LICENSE ─────────────────────────────────────────────
    documents = content.get('documents', {})
    if not isinstance(documents, dict):
        documents = {}
    if any(documents.values()):
        add_section_heading(doc, 'Documents & License')
        if documents.get('coc'):
            add_body_text(doc, f"Certificate of Competency: {documents['coc']}")
        if documents.get('cdc'):
            add_body_text(doc, f"CDC: {documents['cdc']}")
        if documents.get('indos'):
            add_body_text(doc, f"INDOS: {documents['indos']}")
        if documents.get('passport'):
            add_body_text(doc, f"Passport: {documents['passport']}")

    # ── LANGUAGES ─────────────────────────────────────────────────────────
    languages = _as_list(content.get('languages', []))
    if languages:
        add_section_heading(doc, 'Languages')
        add_body_text(doc, '  •  '.join(str(l) for l in languages))

    # ── Save File ─────────────────────────────────────────────────────
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    output_path = os.path.join(OUTPUT_DIR, f'{job_id}_ats_resume.docx')
    doc.save(output_path)
    print(f'✓ DOCX saved: {output_path}')
    return output_path